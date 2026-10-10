"""
Service layer for schedule operations and CP-SAT solver orchestration.
"""

import io
from typing import Optional
import uuid
import datetime
import calendar
from collections import defaultdict
import openpyxl
from openpyxl.styles import PatternFill, Font
from sqlmodel import Session
from ortools.sat.python import cp_model
from sqlmodel import Session

from src.department import repository as department_repository
from src.models import ScheduleConfig
from src.scheduler import ShiftScheduler
from src.schedule.models import ScheduleDraft
from src.schedule.excel_parser import parse_excel_schedule_workbook
from src.schedule import repository as schedule_repository
from src.shift import repository as shift_repository

from src.doctor import repository as doctor_repository
from src.doctor.models import Doctor as DoctorModel
from src.position import repository as position_repository
from src.position.models import Position as PositionModel
from src.shift.models import Shift as ShiftModel, ShiftAssignment as ShiftAssignmentModel
from src.shift.schemas import ShiftCreate


def generate_schedule_from_excel(
    session: Session,
    file_bytes: bytes,
    department_id: uuid.UUID,
    target_month: str,
    source_filename: str,
    config_override: Optional[ScheduleConfig] = None,
) -> ScheduleDraft:
    """
    Generates an optimized monthly schedule from an uploaded Excel workbook.

    Parses the spreadsheet into domain models, executes the Google OR-Tools CP-SAT
    constraint solver to balance duties across doctors and teams, extracts the
    resulting shift assignments, and persists the draft in the database.

    Args:
        session: Active database session.
        file_bytes: Raw binary bytes of the uploaded .xlsx spreadsheet.
        department_id: UUID of the target hospital department.
        target_month: Target month in "YYYY-MM" format (e.g., "2026-11").
        source_filename: Name of the uploaded file for audit and display tracking.
        config_override: Optional custom solver weights and constraint limits.

    Returns:
        The newly created and persisted ScheduleDraft record with its assignments.

    Raises:
        ValueError: If the department does not exist, the spreadsheet format is invalid,
                    or the solver determines the constraints are infeasible.
    """
    department = department_repository.get_department_by_id_global(
        session=session,
        department_id=department_id,
    )

    if not department or department.is_deleted:
        raise ValueError(f"Department not found: {department_id}")

    dept_domain = parse_excel_schedule_workbook(
        file_bytes, department.name, target_month
    )
    if config_override:
        dept_domain.config = config_override

    year, month = map(int, target_month.split("-"))
    scheduler = ShiftScheduler(department=dept_domain)
    status = scheduler.create_schedule(month=month, year=year)
    status_name = scheduler.solver.status_name(status)

    if status not in (cp_model.FEASIBLE, cp_model.OPTIMAL):
        raise ValueError(
            f"Unable to generate schedule. Constraints are infeasible (status: {status_name})."
        )

    assignments = []
    for (day_idx, position, shift, doc), var in scheduler.shift_assignments.items():
        if scheduler.solver.value(var) == 1:
            date = scheduler.dates[day_idx]
            assignments.append({
                "date": date.isoformat(),
                "day_name": date.strftime("%A"),
                "doctor_name": doc.name,
                "doctor_email": doc.email,
                "position": position.name,
                "shift": shift.name,
            })

    assignments.sort(key=lambda a: (a["date"], a["position"], a["shift"]))

    unavailabilities = {
        doc.name: sorted([d.day for d in doc.unavailability])
        for doc in dept_domain.doctors
        if doc.unavailability
    }

    draft = ScheduleDraft(
        department_id=department_id,
        target_month=target_month,
        source_filename=source_filename,
        total_duties=len(assignments),
        solver_status=status_name,
        status="draft",
        assignments=assignments,
        unavailabilities=unavailabilities,
    )

    return schedule_repository.create_schedule_draft(session, draft)


def _populate_doctor_day_cell(
    cell,
    day: int,
    doc_unavail_days: set[int],
    shifts: list[str],
    is_weekend: bool,
    unavailability_fill: PatternFill,
    weekend_fill: PatternFill,
) -> int:
    """Populate a single doctor day cell with fill or text, returning duty count."""
    if day in doc_unavail_days:
        cell.fill = unavailability_fill
        return 0

    if shifts:
        cell.value = ", ".join(shifts)

    if is_weekend:
        cell.fill = weekend_fill

    return len(shifts)


def export_schedule_draft_to_excel(draft: ScheduleDraft) -> bytes:
    """Exports a ScheduleDraft model's assignments into a formatted Excel (.xlsx) file in bytes."""
    # 1. Parse target month
    try:
        parts = draft.target_month.split("-")
        year, month = int(parts[0]), int(parts[1])
    except (ValueError, IndexError) as exc:
        raise ValueError(
            f"Invalid target_month format '{draft.target_month}'. Expected YYYY-MM."
        ) from exc

    if month < 1 or month > 12:
        raise ValueError(
            f"Invalid target_month format '{draft.target_month}'. Expected YYYY-MM."
        )

    days_in_month = calendar.monthrange(year, month)[1]

    # 2. Workbook styling setup
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Schedule {draft.target_month}"

    unavailability_fill = PatternFill(
        start_color="000000", end_color="000000", fill_type="solid"
    )
    weekend_fill = PatternFill(
        start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"
    )
    header_font = Font(bold=True)

    header_cell_1 = ws.cell(row=2, column=1, value="Doctor \\ Date")
    header_cell_1.font = header_font
    header_cell_total = ws.cell(
        row=2, column=days_in_month + 2, value="Total Duties")
    header_cell_total.font = header_font

    # 3. Header rows (1, 2)
    for day_idx in range(1, days_in_month + 1):
        date = datetime.date(year, month, day_idx)
        weekday = date.strftime("%a")

        # Row 1 (weekday name)
        cell = ws.cell(row=1, column=day_idx + 1, value=weekday)
        if date.weekday() >= 5:
            cell.fill = weekend_fill

        # Row 2 (day number)
        cell = ws.cell(row=2, column=day_idx + 1, value=date.day)
        if date.weekday() >= 5:
            cell.fill = weekend_fill

    # 4. If no assignments, save and return empty grid
    if not draft.assignments:
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    # 5. Group assignments by (doctor_name, day_num)
    doctor_day_shifts = defaultdict(list)
    doctor_names = set()

    for assignment in draft.assignments:
        doc_name = assignment["doctor_name"]
        doctor_names.add(doc_name)

        day_num = int(assignment["date"].split("-")[2])
        doctor_day_shifts[(doc_name, day_num)].append(
            assignment.get("shift", "Duty")
        )
    sorted_doctors = sorted(list(doctor_names))

    unavailability_dict = draft.unavailabilities or {}

    # 6. Populate doctor rows
    for row_idx, doc_name in enumerate(sorted_doctors, start=3):
        ws.cell(row=row_idx, column=1, value=doc_name)
        doc_unavail_days = set(unavailability_dict.get(doc_name, []))
        total_duties = 0

        for day in range(1, days_in_month + 1):
            cell = ws.cell(row=row_idx, column=day + 1)
            is_weekend = datetime.date(year, month, day).weekday() >= 5
            shifts = doctor_day_shifts.get((doc_name, day), [])

            total_duties += _populate_doctor_day_cell(
                cell=cell,
                day=day,
                doc_unavail_days=doc_unavail_days,
                shifts=shifts,
                is_weekend=is_weekend,
                unavailability_fill=unavailability_fill,
                weekend_fill=weekend_fill,
            )

        ws.cell(row=row_idx, column=days_in_month + 2, value=total_duties)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _compute_next_month_str(year: int, month: int) -> str:
    """Computes the subsequent calendar month string in YYYY-MM format."""
    if month == 12:
        return f"{year + 1:04d}-01"
    return f"{year:04d}-{month + 1:02d}"


def resolve_next_target_month(
    session: Session,
    department_id: uuid.UUID,
) -> dict[str, str | None]:
    """
    Resolves the target month for schedule generation based on the department's
    latest active ShiftAssignment date. Falls back to next calendar month from today
    if no assignments exist.
    """
    latest_date = shift_repository.get_latest_shift_assignment_date_for_department(
        session=session,
        department_id=department_id,
    )

    if latest_date is not None:
        return {
            "next_target_month": _compute_next_month_str(latest_date.year, latest_date.month),
            "last_published_month": f"{latest_date.year:04d}-{latest_date.month:02d}",
        }

    today = datetime.date.today()
    return {
        "next_target_month": _compute_next_month_str(today.year, today.month),
        "last_published_month": None,
    }


def publish_schedule_draft(
    session: Session,
    draft_id: uuid.UUID,
    department_id: uuid.UUID,
) -> ScheduleDraft:
    """Finalizes a schedule draft and publishes it to the db."""
    draft = schedule_repository.get_schedule_draft_by_id(session, draft_id)
    if draft is None or draft.is_deleted or draft.department_id != department_id:
        raise ValueError(f"Schedule draft not found: {draft_id}.")

    if draft.status == "published":
        raise ValueError(f"Schedule draft {draft_id} is already published.")

    year, month = map(int, draft.target_month.split("-"))

    # 1. Soft-delete prior month assignments (uncommitted flush)
    shift_repository.soft_delete_shift_assignments_for_department_month(
        session=session,
        department_id=department_id,
        year=year,
        month=month,
        commit=False
    )

    # 2. Pre-fetch existing entities to eliminate N+1 database queries
    doc_cache: dict[str, DoctorModel] = {
        d.name: d for d in doctor_repository.get_active_doctors_for_department(session, department_id)
    }
    pos_cache: dict[str, PositionModel] = {
        p.name: p for p in position_repository.get_active_positions_for_department(session, department_id)
    }
    shift_cache: dict[tuple[uuid.UUID, str], ShiftModel] = {
        (s.position_id, s.name): s for s in shift_repository.get_active_shifts_for_department(session, department_id)
    }

    new_assignments: list[ShiftAssignmentModel] = []
    for item in draft.assignments:
        pos_name = item.get("position", "General")
        shift_name = item.get("shift", "Duty")
        doc_name = item.get("doctor_name") or item.get(
            "doctor") or "Unknown Doctor"
        assignment_date = datetime.date.fromisoformat(item["date"])

        # Resolve or auto-provision Position
        if pos_name not in pos_cache:
            pos = position_repository.stage_position(
                session=session,
                position_name=pos_name,
                duty_days=[1, 2, 3, 4, 5, 6, 7],
                department_id=department_id
            )
            pos_cache[pos_name] = pos
        else:
            pos = pos_cache[pos_name]

        # Resolve or auto-provision Shift
        shift_key = (pos.id, shift_name)
        if shift_key not in shift_cache:
            shift = shift_repository.stage_shift(
                session=session,
                shift_data=ShiftCreate(
                    name=shift_name,
                    doctors_per_shift=1,
                    grants_day_off=True,
                    position_id=pos.id
                )
            )
            shift_cache[shift_key] = shift
        else:
            shift = shift_cache[shift_key]

        # Resolve or auto-provision Doctor
        if doc_name not in doc_cache:
            doc = doctor_repository.stage_doctor(
                session=session,
                name=doc_name,
                department_id=department_id,
            )
            doc_cache[doc_name] = doc
        else:
            doc = doc_cache[doc_name]

        new_assignments.append(
            ShiftAssignmentModel(
                doctor_id=doc.id,
                shift_id=shift.id,
                date=assignment_date
            )
        )

    # 3. Bulk insert all assignments (uncommitted flush)
    shift_repository.bulk_create_shift_assignments(
        session=session,
        assignments=new_assignments,
        commit=False,
    )

    draft = schedule_repository.update_schedule_draft_status(
        session=session,
        draft=draft,
        status="published",
        commit=True
    )

    return draft
