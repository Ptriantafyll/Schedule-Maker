"""Unit tests for the Excel schedule workbook parser."""

import io
import datetime
import pytest
from openpyxl import Workbook
from src.schedule.excel_parser import parse_excel_schedule_workbook


def _create_workbook_bytes(
    sheets_or_rows: dict[str, list[dict]] | list[dict],
    sheet_name: str = "Doctors",
) -> bytes:
    """Helper to generate an in-memory .xlsx file from a list of dicts or multi-sheet dict."""
    if isinstance(sheets_or_rows, list):
        sheets = {sheet_name: sheets_or_rows}
    else:
        sheets = sheets_or_rows

    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(title=name)
        ws.title = name
        first = False
        if rows:
            headers = list(rows[0].keys())
            ws.append(headers)
            for row in rows:
                ws.append([row.get(h) for h in headers])

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def test_parse_valid_excel_with_all_teams_defined():
    """Verify parsing a valid workbook where all doctors have team assignments."""
    rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "Alpha",
            "Unavailability": "1, 2, 15",
            "Pre-Assignments": "5:Night",
            "Max Duties": 6,
        },
        {
            "Name": "Dr. Allison Cameron",
            "Email": "cameron@hospital.org",
            "Position": "ER",
            "Team": "Alpha",
            "Unavailability": "3, 20",
            "Pre-Assignments": "",
            "Max Duties": None,
        },
        {
            "Name": "Dr. Eric Foreman",
            "Email": "foreman@hospital.org",
            "Position": "ICU",
            "Team": "Beta",
            "Unavailability": "",
            "Pre-Assignments": "10:Morning",
            "Max Duties": 5,
        },
    ]

    file_bytes = _create_workbook_bytes(rows)
    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Emergency",
        target_month="2026-11",
    )

    assert dept.name == "Emergency"
    assert len(dept.doctors) == 3

    # Verify Teams
    team_names = {t.name for t in dept.teams}
    assert team_names == {"Alpha", "Beta"}
    alpha_team = next(t for t in dept.teams if t.name == "Alpha")
    assert len(alpha_team.doctors) == 2

    # Verify Doctor 1 details
    house = next(d for d in dept.doctors if d.name == "Dr. Gregory House")
    assert house.email == "house@hospital.org"
    assert house.unavailability == {
        datetime.date(2026, 11, 1),
        datetime.date(2026, 11, 2),
        datetime.date(2026, 11, 15),
    }
    assert len(house.pre_assignments) == 1
    pre_date, pre_shift = house.pre_assignments[0]
    assert pre_date == datetime.date(2026, 11, 5)
    assert pre_shift.name == "Night"

    # Verify Positions
    pos_names = {p.name for p in dept.positions}
    assert pos_names == {"ER", "ICU"}


def test_parse_valid_excel_with_all_teams_blank():
    """Verify that when all team values are blank, a single unified team is created."""
    rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "",
            "Unavailability": "1, 2",
            "Pre-Assignments": "",
        },
        {
            "Name": "Dr. Allison Cameron",
            "Email": "cameron@hospital.org",
            "Position": "ER",
            "Team": None,
            "Unavailability": "",
            "Pre-Assignments": "",
        },
    ]

    file_bytes = _create_workbook_bytes(rows)
    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Pediatrics",
        target_month="2026-11",
    )

    assert len(dept.teams) == 1
    assert dept.teams[0].name == "Team Pediatrics"
    assert len(dept.teams[0].doctors) == 2


def test_parse_mixed_teams_raises_validation_error():
    """Verify that having some doctors with teams and some without raises a ValueError."""
    rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "Alpha",
            "Unavailability": "",
            "Pre-Assignments": "",
        },
        {
            "Name": "Dr. Allison Cameron",
            "Email": "cameron@hospital.org",
            "Position": "ER",
            "Team": "",  # Blank while another has Alpha!
            "Unavailability": "",
            "Pre-Assignments": "",
        },
    ]

    file_bytes = _create_workbook_bytes(rows)
    with pytest.raises(ValueError, match="Team"):
        parse_excel_schedule_workbook(
            file_bytes=file_bytes,
            department_name="Emergency",
            target_month="2026-11",
        )


def test_missing_doctors_sheet_raises_error():
    """Verify that a workbook without a 'Doctors' sheet raises a ValueError."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ScheduleData"
    ws.append(["Name", "Email", "Position"])
    output = io.BytesIO()
    wb.save(output)

    with pytest.raises(ValueError, match="Doctors"):
        parse_excel_schedule_workbook(
            file_bytes=output.getvalue(),
            department_name="Emergency",
            target_month="2026-11",
        )


def test_missing_required_column_raises_error():
    """Verify that missing required columns like Position raises a ValueError."""
    rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            # Missing Position!
            "Team": "Alpha",
        }
    ]

    file_bytes = _create_workbook_bytes(rows)
    with pytest.raises(ValueError, match="Position"):
        parse_excel_schedule_workbook(
            file_bytes=file_bytes,
            department_name="Emergency",
            target_month="2026-11",
        )


def test_out_of_range_unavailability_day_raises_error():
    """Verify that a day number out of range for the target month raises a ValueError."""
    rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Unavailability": "35",  # Invalid day for November!
        }
    ]

    file_bytes = _create_workbook_bytes(rows)
    with pytest.raises(ValueError, match="35"):
        parse_excel_schedule_workbook(
            file_bytes=file_bytes,
            department_name="Emergency",
            target_month="2026-11",
        )


# --- 2-Sheet Tests (Doctors + Shifts) ---


def test_parse_excel_with_explicit_shifts_sheet():
    """Verify parsing when both Doctors and Shifts sheets are provided."""
    doctor_rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "Team 1",
            "Unavailability": "",
            "Pre-Assignments": "5:Night",
        },
        {
            "Name": "Dr. Allison Cameron",
            "Email": "cameron@hospital.org",
            "Position": "ER",
            "Team": "Team 1",
            "Unavailability": "",
            "Pre-Assignments": "",
        },
        {
            "Name": "Dr. Eric Foreman",
            "Email": "foreman@hospital.org",
            "Position": "ICU",
            "Team": "Team 2",
            "Unavailability": "",
            "Pre-Assignments": "",
        },
    ]

    shift_rows = [
        {
            "Position": "ER",
            "Shift Name": "Night",
            "Doctors Per Shift": 2,
            "Duty Days": "Mon-Fri",
            "Grants Day Off": "Yes",
        },
        {
            "Position": "ICU",
            "Shift Name": "On-Call",
            "Doctors Per Shift": 1,
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "No",
        },
    ]

    file_bytes = _create_workbook_bytes({
        "Doctors": doctor_rows,
        "Shifts": shift_rows,
    })

    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Emergency",
        target_month="2026-11",
    )

    # Check positions
    er_pos = next(p for p in dept.positions if p.name == "ER")
    icu_pos = next(p for p in dept.positions if p.name == "ICU")

    # ER checks: Mon-Fri = {0, 1, 2, 3, 4}
    assert er_pos.duty_days == {0, 1, 2, 3, 4}
    assert len(er_pos.shifts) == 1
    assert er_pos.shifts[0].name == "Night"
    assert er_pos.shifts[0].doctors_per_shift == 2
    assert er_pos.shifts[0].grants_day_off is True
    assert {d.name for d in er_pos.eligible_doctors} == {"Dr. Gregory House", "Dr. Allison Cameron"}

    # ICU checks: Mon-Sun = {0, 1, 2, 3, 4, 5, 6}
    assert icu_pos.duty_days == {0, 1, 2, 3, 4, 5, 6}
    assert len(icu_pos.shifts) == 1
    assert icu_pos.shifts[0].name == "On-Call"
    assert icu_pos.shifts[0].doctors_per_shift == 1
    assert icu_pos.shifts[0].grants_day_off is False
    assert {d.name for d in icu_pos.eligible_doctors} == {"Dr. Eric Foreman"}


def test_parse_fallback_when_shifts_sheet_omitted():
    """Verify that omitting the Shifts sheet defaults to 1 doctor/day Mon-Sun for each position."""
    doctor_rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "Alpha",
            "Unavailability": "",
            "Pre-Assignments": "",
        },
        {
            "Name": "Dr. Eric Foreman",
            "Email": "foreman@hospital.org",
            "Position": "ICU",
            "Team": "Alpha",
            "Unavailability": "",
            "Pre-Assignments": "",
        },
    ]

    file_bytes = _create_workbook_bytes({"Doctors": doctor_rows})

    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Emergency",
        target_month="2026-11",
    )

    for pos in dept.positions:
        assert pos.duty_days == {0, 1, 2, 3, 4, 5, 6}
        assert len(pos.shifts) == 1
        assert pos.shifts[0].doctors_per_shift == 1
        assert pos.shifts[0].grants_day_off is True


def test_shifts_sheet_unknown_position_raises_error():
    """Verify that a shift referencing a non-existent position raises a ValueError."""
    doctor_rows = [
        {
            "Name": "Dr. House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "Alpha",
        }
    ]
    shift_rows = [
        {
            "Position": "Surgery",  # Not in Doctors sheet!
            "Shift Name": "On-Call",
            "Doctors Per Shift": 1,
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "No",
        }
    ]

    file_bytes = _create_workbook_bytes({
        "Doctors": doctor_rows,
        "Shifts": shift_rows,
    })

    with pytest.raises(ValueError, match="Surgery"):
        parse_excel_schedule_workbook(
            file_bytes=file_bytes,
            department_name="Emergency",
            target_month="2026-11",
        )


def test_shifts_sheet_invalid_doctors_per_shift_raises_error():
    """Verify that a non-positive doctors_per_shift raises a ValueError."""
    doctor_rows = [
        {
            "Name": "Dr. House",
            "Email": "house@hospital.org",
            "Position": "ER",
            "Team": "Alpha",
        }
    ]
    shift_rows = [
        {
            "Position": "ER",
            "Shift Name": "Night",
            "Doctors Per Shift": 0,  # Invalid!
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "No",
        }
    ]

    file_bytes = _create_workbook_bytes({
        "Doctors": doctor_rows,
        "Shifts": shift_rows,
    })

    with pytest.raises(ValueError, match="Doctors Per Shift"):
        parse_excel_schedule_workbook(
            file_bytes=file_bytes,
            department_name="Emergency",
            target_month="2026-11",
        )


def test_parse_doctor_with_multiple_positions_default_shifts():
    """Verify that a doctor with comma-separated positions is added to both positions."""
    doctor_rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER, ICU",
        },
        {
            "Name": "Dr. James Wilson",
            "Email": "wilson@hospital.org",
            "Position": "ER",
        },
    ]

    file_bytes = _create_workbook_bytes({"Doctors": doctor_rows})

    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Medicine",
        target_month="2026-11",
    )

    pos_dict = {p.name: p for p in dept.positions}
    assert "ER" in pos_dict
    assert "ICU" in pos_dict

    # Check eligible doctors for ER
    er_names = [d.name for d in pos_dict["ER"].eligible_doctors]
    assert "Dr. Gregory House" in er_names
    assert "Dr. James Wilson" in er_names

    # Check eligible doctors for ICU
    icu_names = [d.name for d in pos_dict["ICU"].eligible_doctors]
    assert "Dr. Gregory House" in icu_names
    assert "Dr. James Wilson" not in icu_names

    # Verify it is the exact same Doctor object reference
    house_er = next(d for d in pos_dict["ER"].eligible_doctors if d.name == "Dr. Gregory House")
    house_icu = next(d for d in pos_dict["ICU"].eligible_doctors if d.name == "Dr. Gregory House")
    assert house_er is house_icu


def test_parse_doctor_with_multiple_positions_explicit_shifts_sheet():
    """Verify that multiple positions work when a Shifts sheet explicitly defines them."""
    doctor_rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER, ICU",
        },
    ]
    shift_rows = [
        {
            "Position": "ER",
            "Shift Name": "Day",
            "Doctors Per Shift": 1,
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "No",
        },
        {
            "Position": "ICU",
            "Shift Name": "Night",
            "Doctors Per Shift": 1,
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "Yes",
        },
    ]

    file_bytes = _create_workbook_bytes({
        "Doctors": doctor_rows,
        "Shifts": shift_rows,
    })

    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Medicine",
        target_month="2026-11",
    )

    pos_dict = {p.name: p for p in dept.positions}
    assert "ER" in pos_dict
    assert "ICU" in pos_dict
    assert len(pos_dict["ER"].shifts) == 1
    assert pos_dict["ER"].shifts[0].name == "Day"
    assert len(pos_dict["ICU"].shifts) == 1
    assert pos_dict["ICU"].shifts[0].name == "Night"

    assert pos_dict["ER"].eligible_doctors[0].name == "Dr. Gregory House"
    assert pos_dict["ICU"].eligible_doctors[0].name == "Dr. Gregory House"
    assert pos_dict["ER"].eligible_doctors[0] is pos_dict["ICU"].eligible_doctors[0]


def test_parse_doctor_multiple_positions_pre_assignments_cross_position():
    """Verify that pre-assignments can reference shifts across any of the doctor's eligible positions."""
    doctor_rows = [
        {
            "Name": "Dr. Gregory House",
            "Email": "house@hospital.org",
            "Position": "ER, ICU",
            "Pre-Assignments": "5:Day, 10:Night",
        },
    ]
    shift_rows = [
        {
            "Position": "ER",
            "Shift Name": "Day",
            "Doctors Per Shift": 1,
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "No",
        },
        {
            "Position": "ICU",
            "Shift Name": "Night",
            "Doctors Per Shift": 1,
            "Duty Days": "Mon-Sun",
            "Grants Day Off": "Yes",
        },
    ]

    file_bytes = _create_workbook_bytes({
        "Doctors": doctor_rows,
        "Shifts": shift_rows,
    })

    dept = parse_excel_schedule_workbook(
        file_bytes=file_bytes,
        department_name="Medicine",
        target_month="2026-11",
    )

    house = dept.positions[0].eligible_doctors[0]
    assert len(house.pre_assignments) == 2

    day_assignment = next(pa for pa in house.pre_assignments if pa[0] == datetime.date(2026, 11, 5))
    night_assignment = next(pa for pa in house.pre_assignments if pa[0] == datetime.date(2026, 11, 10))

    assert day_assignment[1].name == "Day"
    assert night_assignment[1].name == "Night"
