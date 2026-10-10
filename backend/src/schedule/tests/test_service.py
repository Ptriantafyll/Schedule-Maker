"""Unit tests for the schedule service layer."""

import io
import uuid
import datetime
import openpyxl
import pytest
from openpyxl import Workbook
from src.department.models import Department as DepartmentModel
from src.models import ScheduleConfig
from src.schedule.models import ScheduleDraft
from src.schedule.service import (
    generate_schedule_from_excel,
    export_schedule_draft_to_excel,
)


def _create_test_workbook_bytes(doctor_rows: list[dict]) -> bytes:
    """Helper to generate an in-memory .xlsx file with a Doctors sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Doctors"
    if doctor_rows:
        headers = list(doctor_rows[0].keys())
        ws.append(headers)
        for row in doctor_rows:
            ws.append([row.get(h) for h in headers])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def test_generate_schedule_from_excel_success(session):
    """Verify that a valid Excel workbook generates and persists a schedule draft."""
    # 1. Create Department in DB
    dept = DepartmentModel(name="Emergency Medicine", code="EMERG")
    session.add(dept)
    session.commit()
    session.refresh(dept)

    # 2. 6 doctors for 30 days in November ensures feasibility with standard limits
    doctor_rows = [
        {"Name": "Dr. Gregory House", "Email": "house@hospital.org", "Position": "ER", "Team": "Alpha", "Unavailability": "14"},
        {"Name": "Dr. Allison Cameron", "Email": "cameron@hospital.org", "Position": "ER", "Team": "Alpha"},
        {"Name": "Dr. Eric Foreman", "Email": "foreman@hospital.org", "Position": "ER", "Team": "Alpha"},
        {"Name": "Dr. Robert Chase", "Email": "chase@hospital.org", "Position": "ER", "Team": "Beta"},
        {"Name": "Dr. James Wilson", "Email": "wilson@hospital.org", "Position": "ER", "Team": "Beta"},
        {"Name": "Dr. Lisa Cuddy", "Email": "cuddy@hospital.org", "Position": "ER", "Team": "Beta"},
    ]
    file_bytes = _create_test_workbook_bytes(doctor_rows)

    # 3. Call Service
    draft = generate_schedule_from_excel(
        session=session,
        file_bytes=file_bytes,
        department_id=dept.id,
        target_month="2026-11",
        source_filename="november_roster.xlsx",
    )

    # 4. Assertions
    assert draft.id is not None
    assert draft.department_id == dept.id
    assert draft.target_month == "2026-11"
    assert draft.source_filename == "november_roster.xlsx"
    assert draft.solver_status in ("OPTIMAL", "FEASIBLE")
    assert draft.status == "draft"
    assert draft.total_duties == 30  # 30 days in November
    assert len(draft.assignments) == 30
    assert draft.unavailabilities == {"Dr. Gregory House": [14]}

    # Verify assignment item structure
    first_assignment = draft.assignments[0]
    assert "date" in first_assignment
    assert "day_name" in first_assignment
    assert "doctor_name" in first_assignment
    assert "doctor_email" in first_assignment
    assert "position" in first_assignment
    assert "shift" in first_assignment


def test_generate_schedule_nonexistent_department_raises_error(session):
    """Verify that attempting to generate for a missing department raises ValueError."""
    fake_dept_id = uuid.uuid4()
    doctor_rows = [
        {"Name": "Dr. Gregory House", "Email": "house@hospital.org", "Position": "ER"}
    ]
    file_bytes = _create_test_workbook_bytes(doctor_rows)

    with pytest.raises(ValueError, match="Department not found"):
        generate_schedule_from_excel(
            session=session,
            file_bytes=file_bytes,
            department_id=fake_dept_id,
            target_month="2026-11",
            source_filename="roster.xlsx",
        )


def test_generate_schedule_soft_deleted_department_raises_error(session):
    """Verify that attempting to generate for a soft-deleted department raises ValueError."""
    dept = DepartmentModel(name="Archived Clinic", code="ARCH", is_deleted=True)
    session.add(dept)
    session.commit()
    session.refresh(dept)

    doctor_rows = [
        {"Name": "Dr. Gregory House", "Email": "house@hospital.org", "Position": "ER"}
    ]
    file_bytes = _create_test_workbook_bytes(doctor_rows)

    with pytest.raises(ValueError, match="Department not found"):
        generate_schedule_from_excel(
            session=session,
            file_bytes=file_bytes,
            department_id=dept.id,
            target_month="2026-11",
            source_filename="roster.xlsx",
        )


def test_generate_schedule_infeasible_constraints_raises_error(session):
    """Verify that impossible constraints (e.g. 1 doctor for 30 shifts) raises ValueError."""
    dept = DepartmentModel(name="Understaffed Dept", code="UNDER")
    session.add(dept)
    session.commit()
    session.refresh(dept)

    # 1 single doctor cannot cover 30 days due to max_duties and no-consecutive constraints
    doctor_rows = [
        {"Name": "Dr. Lone Ranger", "Email": "lone@hospital.org", "Position": "ER", "Team": "Solo"}
    ]
    file_bytes = _create_test_workbook_bytes(doctor_rows)

    with pytest.raises(ValueError, match="infeasible"):
        generate_schedule_from_excel(
            session=session,
            file_bytes=file_bytes,
            department_id=dept.id,
            target_month="2026-11",
            source_filename="impossible.xlsx",
        )


def test_generate_schedule_corrupted_excel_file_raises_error(session):
    """Verify that passing corrupted/non-Excel binary bytes raises ValueError."""
    dept = DepartmentModel(name="Emergency", code="EMERG")
    session.add(dept)
    session.commit()
    session.refresh(dept)

    corrupted_bytes = b"This is plain text and definitely not a valid xlsx zip archive"

    with pytest.raises(ValueError, match="Invalid or corrupted Excel file"):
        generate_schedule_from_excel(
            session=session,
            file_bytes=corrupted_bytes,
            department_id=dept.id,
            target_month="2026-11",
            source_filename="corrupted.xlsx",
        )


def test_generate_schedule_with_custom_config_override(session):
    """Verify that custom ScheduleConfig overrides are applied and generate successfully."""
    dept = DepartmentModel(name="Flexible Dept", code="FLEX")
    session.add(dept)
    session.commit()
    session.refresh(dept)

    doctor_rows = [
        {"Name": "Dr. A", "Email": "a@hospital.org", "Position": "ER"},
        {"Name": "Dr. B", "Email": "b@hospital.org", "Position": "ER"},
        {"Name": "Dr. C", "Email": "c@hospital.org", "Position": "ER"},
        {"Name": "Dr. D", "Email": "d@hospital.org", "Position": "ER"},
        {"Name": "Dr. E", "Email": "e@hospital.org", "Position": "ER"},
        {"Name": "Dr. F", "Email": "f@hospital.org", "Position": "ER"},
    ]
    file_bytes = _create_test_workbook_bytes(doctor_rows)

    custom_config = ScheduleConfig(
        solver_time_limit=15,
        max_duties_per_month=7,
        w_every_other_penalty=10,
    )

    draft = generate_schedule_from_excel(
        session=session,
        file_bytes=file_bytes,
        department_id=dept.id,
        target_month="2026-11",
        source_filename="custom_config.xlsx",
        config_override=custom_config,
    )

    assert draft.id is not None
    assert draft.solver_status in ("OPTIMAL", "FEASIBLE")
    assert len(draft.assignments) == 30


def test_export_schedule_draft_to_excel_success():
    """Verify that export_schedule_draft_to_excel produces a correctly formatted Excel grid."""
    draft = ScheduleDraft(
        department_id=uuid.uuid4(),
        target_month="2026-11",
        source_filename="november_roster.xlsx",
        total_duties=3,
        solver_status="OPTIMAL",
        assignments=[
            {
                "date": "2026-11-01",
                "day_name": "Sunday",
                "doctor_name": "Dr. Gregory House",
                "doctor_email": "house@hospital.org",
                "position": "ER",
                "shift": "Duty",
            },
            {
                "date": "2026-11-02",
                "day_name": "Monday",
                "doctor_name": "Dr. Allison Cameron",
                "doctor_email": "cameron@hospital.org",
                "position": "ER",
                "shift": "Duty",
            },
            {
                "date": "2026-11-03",
                "day_name": "Tuesday",
                "doctor_name": "Dr. Gregory House",
                "doctor_email": "house@hospital.org",
                "position": "ER",
                "shift": "Duty",
            },
        ],
        unavailabilities={"Dr. Gregory House": [14]},
    )

    file_bytes = export_schedule_draft_to_excel(draft)
    assert isinstance(file_bytes, bytes)
    assert len(file_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    assert ws is not None

    # Check header row 1 (weekdays)
    assert ws.cell(row=1, column=2).value == "Sun"
    assert ws.cell(row=1, column=3).value == "Mon"
    assert ws.cell(row=1, column=4).value == "Tue"

    # Check header row 2 (day numbers & labels)
    assert ws.cell(row=2, column=1).value == "Doctor \\ Date"
    assert ws.cell(row=2, column=2).value == 1
    assert ws.cell(row=2, column=3).value == 2
    assert ws.cell(row=2, column=31).value == 30  # November has 30 days
    assert ws.cell(row=2, column=32).value == "Total Duties"

    # Check doctor names are present in column 1
    doctor_names = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
    assert "Dr. Gregory House" in doctor_names
    assert "Dr. Allison Cameron" in doctor_names

    # Locate Dr. House row
    house_row = 3 if ws.cell(row=3, column=1).value == "Dr. Gregory House" else 4
    cameron_row = 4 if house_row == 3 else 3

    # House: Nov 1 = "Duty", Nov 2 = None, Nov 3 = "Duty", Total = 2
    assert ws.cell(row=house_row, column=2).value == "Duty"
    assert ws.cell(row=house_row, column=3).value is None
    assert ws.cell(row=house_row, column=4).value == "Duty"
    assert ws.cell(row=house_row, column=32).value == 2

    # Cameron: Nov 2 = "Duty", Total = 1
    assert ws.cell(row=cameron_row, column=3).value == "Duty"
    assert ws.cell(row=cameron_row, column=32).value == 1

    # Verify weekend column has fill
    weekend_cell = ws.cell(row=1, column=2)  # Sun Nov 1
    assert weekend_cell.fill.fill_type is not None

    # Verify Dr. House unavailable cell (Nov 14 -> Col 15) has black fill
    house_unavail_cell = ws.cell(row=house_row, column=15)
    assert house_unavail_cell.fill.fill_type == "solid"
    assert str(house_unavail_cell.fill.start_color.rgb).endswith("000000")



def test_export_schedule_draft_empty_assignments():
    """Verify that exporting a draft with no assignments produces valid headers and no crashes."""
    draft = ScheduleDraft(
        department_id=uuid.uuid4(),
        target_month="2026-11",
        source_filename="empty.xlsx",
        total_duties=0,
        solver_status="OPTIMAL",
        assignments=[],
    )

    file_bytes = export_schedule_draft_to_excel(draft)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active

    assert ws.cell(row=2, column=1).value == "Doctor \\ Date"
    assert ws.cell(row=2, column=32).value == "Total Duties"
    assert ws.max_row == 2


def test_export_schedule_draft_multiple_shifts_same_day():
    """Verify that multiple shifts for the same doctor on the same day are joined with commas."""
    draft = ScheduleDraft(
        department_id=uuid.uuid4(),
        target_month="2026-11",
        source_filename="multi.xlsx",
        total_duties=2,
        solver_status="OPTIMAL",
        assignments=[
            {
                "date": "2026-11-05",
                "day_name": "Thursday",
                "doctor_name": "Dr. Gregory House",
                "doctor_email": "house@hospital.org",
                "position": "ER",
                "shift": "Morning",
            },
            {
                "date": "2026-11-05",
                "day_name": "Thursday",
                "doctor_name": "Dr. Gregory House",
                "doctor_email": "house@hospital.org",
                "position": "ER",
                "shift": "Night",
            },
        ],
    )

    file_bytes = export_schedule_draft_to_excel(draft)
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    ws = wb.active

    # Nov 5 is column 6 (Col 1 is Doctor, Col 2 is Day 1 -> Col 6 is Day 5)
    assert ws.cell(row=3, column=6).value == "Morning, Night"
    assert ws.cell(row=3, column=32).value == 2


def test_export_schedule_draft_invalid_target_month_raises_value_error():
    """Verify that an invalid target_month string raises ValueError."""
    draft = ScheduleDraft(
        department_id=uuid.uuid4(),
        target_month="invalid-month",
        source_filename="test.xlsx",
        total_duties=0,
        solver_status="OPTIMAL",
        assignments=[],
    )

    with pytest.raises(ValueError) as exc_info:
        export_schedule_draft_to_excel(draft)

    assert "Invalid target_month format" in str(exc_info.value)


def test_resolve_next_target_month_advances_following_month(session):
    """Verify that target month advances to the month following the latest shift assignment."""
    from src.doctor.models import Doctor as DoctorModel
    from src.position.models import Position as PositionModel
    from src.shift.models import Shift as ShiftModel, ShiftAssignment as ShiftAssignmentModel
    from src.schedule.service import resolve_next_target_month

    dept = DepartmentModel(name="Cardiology", code="CARD")
    session.add(dept)
    session.commit()

    pos = PositionModel(name="Cardio Ward", department_id=dept.id)
    session.add(pos)
    session.commit()

    shift = ShiftModel(name="Morning", position_id=pos.id)
    doc = DoctorModel(name="Dr. Adams", department_id=dept.id)
    session.add(shift)
    session.add(doc)
    session.commit()

    # Latest assignment on May 31, 2026
    assignment = ShiftAssignmentModel(
        doctor_id=doc.id,
        shift_id=shift.id,
        date=datetime.date(2026, 5, 31),
    )
    session.add(assignment)
    session.commit()

    result = resolve_next_target_month(session=session, department_id=dept.id)

    assert result["last_published_month"] == "2026-05"
    assert result["next_target_month"] == "2026-06"


def test_resolve_next_target_month_handles_december_year_rollover(session):
    """Verify that target month rolls over from December to January of the following year."""
    from src.doctor.models import Doctor as DoctorModel
    from src.position.models import Position as PositionModel
    from src.shift.models import Shift as ShiftModel, ShiftAssignment as ShiftAssignmentModel
    from src.schedule.service import resolve_next_target_month

    dept = DepartmentModel(name="Neurology", code="NEURO")
    session.add(dept)
    session.commit()

    pos = PositionModel(name="Neuro Ward", department_id=dept.id)
    session.add(pos)
    session.commit()

    shift = ShiftModel(name="Night", position_id=pos.id)
    doc = DoctorModel(name="Dr. Cuddy", department_id=dept.id)
    session.add(shift)
    session.add(doc)
    session.commit()

    # Assignment on Dec 25, 2026 -> should roll to 2027-01
    assignment = ShiftAssignmentModel(
        doctor_id=doc.id,
        shift_id=shift.id,
        date=datetime.date(2026, 12, 25),
    )
    session.add(assignment)
    session.commit()

    result = resolve_next_target_month(session=session, department_id=dept.id)

    assert result["last_published_month"] == "2026-12"
    assert result["next_target_month"] == "2027-01"


def test_resolve_next_target_month_empty_department_falls_back_to_next_calendar_month(session):
    """Verify that an empty department defaults to next calendar month from today with last_published=None."""
    from src.schedule.service import resolve_next_target_month

    dept = DepartmentModel(name="Pediatrics", code="PED")
    session.add(dept)
    session.commit()

    result = resolve_next_target_month(session=session, department_id=dept.id)

    # Compute expected next calendar month relative to date.today()
    today = datetime.date.today()
    expected_year = today.year if today.month < 12 else today.year + 1
    expected_month = today.month + 1 if today.month < 12 else 1
    expected_target = f"{expected_year:04d}-{expected_month:02d}"

    assert result["last_published_month"] is None
    assert result["next_target_month"] == expected_target


def test_resolve_next_target_month_ignores_soft_deleted_assignments(session):
    """Verify that soft-deleted assignments are ignored during target month calculation."""
    from src.doctor.models import Doctor as DoctorModel
    from src.position.models import Position as PositionModel
    from src.shift.models import Shift as ShiftModel, ShiftAssignment as ShiftAssignmentModel
    from src.schedule.service import resolve_next_target_month

    dept = DepartmentModel(name="Orthopedics", code="ORTHO")
    session.add(dept)
    session.commit()

    pos = PositionModel(name="Ortho Ward", department_id=dept.id)
    session.add(pos)
    session.commit()

    shift = ShiftModel(name="Duty", position_id=pos.id)
    doc = DoctorModel(name="Dr. Wilson", department_id=dept.id)
    session.add(shift)
    session.add(doc)
    session.commit()

    deleted_assignment = ShiftAssignmentModel(
        doctor_id=doc.id,
        shift_id=shift.id,
        date=datetime.date(2026, 10, 15),
        is_deleted=True,
    )
    session.add(deleted_assignment)
    session.commit()

    result = resolve_next_target_month(session=session, department_id=dept.id)

    # Soft-deleted assignment should not count as published
    assert result["last_published_month"] is None


def test_publish_schedule_draft_success(session):
    """Verify publishing a draft materializes ShiftAssignments and marks draft published."""
    from src.schedule.service import publish_schedule_draft
    from src.shift.models import ShiftAssignment as ShiftAssignmentModel

    dept = DepartmentModel(name="General Medicine", code="GEN")
    session.add(dept)
    session.commit()

    draft = ScheduleDraft(
        department_id=dept.id,
        target_month="2026-11",
        source_filename="november_roster.xlsx",
        total_duties=2,
        solver_status="OPTIMAL",
        status="draft",
        assignments=[
            {
                "date": "2026-11-01",
                "day_name": "Sunday",
                "doctor_name": "Dr. Gregory House",
                "doctor_email": "house@hospital.org",
                "position": "ER",
                "shift": "Duty",
            },
            {
                "date": "2026-11-02",
                "day_name": "Monday",
                "doctor_name": "Dr. Allison Cameron",
                "doctor_email": "cameron@hospital.org",
                "position": "ER",
                "shift": "Duty",
            },
        ],
    )
    session.add(draft)
    session.commit()

    published_draft = publish_schedule_draft(
        session=session,
        draft_id=draft.id,
        department_id=dept.id,
    )

    assert published_draft.status == "published"

    # Verify ShiftAssignment rows created in database
    assignments_in_db = session.query(ShiftAssignmentModel).filter(
        ShiftAssignmentModel.is_deleted == False
    ).all()
    assert len(assignments_in_db) == 2
    dates = {str(a.date) for a in assignments_in_db}
    assert dates == {"2026-11-01", "2026-11-02"}


def test_publish_schedule_draft_already_published_raises_error(session):
    """Verify attempting to publish an already published draft raises ValueError."""
    from src.schedule.service import publish_schedule_draft

    dept = DepartmentModel(name="Surgery", code="SURG")
    session.add(dept)
    session.commit()

    draft = ScheduleDraft(
        department_id=dept.id,
        target_month="2026-11",
        source_filename="november_roster.xlsx",
        status="published",
        assignments=[],
    )
    session.add(draft)
    session.commit()

    with pytest.raises(ValueError, match="already published"):
        publish_schedule_draft(
            session=session,
            draft_id=draft.id,
            department_id=dept.id,
        )


def test_publish_schedule_draft_nonexistent_or_foreign_raises_error(session):
    """Verify attempting to publish a nonexistent or foreign draft raises ValueError."""
    from src.schedule.service import publish_schedule_draft

    dept_a = DepartmentModel(name="Dept A", code="DA")
    dept_b = DepartmentModel(name="Dept B", code="DB")
    session.add_all([dept_a, dept_b])
    session.commit()

    draft = ScheduleDraft(
        department_id=dept_a.id,
        target_month="2026-11",
        source_filename="test.xlsx",
        status="draft",
        assignments=[],
    )
    session.add(draft)
    session.commit()

    # Nonexistent draft ID
    with pytest.raises(ValueError, match="not found"):
        publish_schedule_draft(
            session=session,
            draft_id=uuid.uuid4(),
            department_id=dept_a.id,
        )

    # Department mismatch (caller is dept_b, draft belongs to dept_a)
    with pytest.raises(ValueError, match="not found"):
        publish_schedule_draft(
            session=session,
            draft_id=draft.id,
            department_id=dept_b.id,
        )


def test_publish_schedule_draft_idempotency_cleans_prior_assignments(session):
    """Verify publishing replaces existing assignments for that month (BL-015)."""
    from src.schedule.service import publish_schedule_draft
    from src.doctor.models import Doctor as DoctorModel
    from src.position.models import Position as PositionModel
    from src.shift.models import Shift as ShiftModel, ShiftAssignment as ShiftAssignmentModel

    dept = DepartmentModel(name="ICU Dept", code="ICU")
    session.add(dept)
    session.commit()

    pos = PositionModel(name="ICU", department_id=dept.id)
    session.add(pos)
    session.commit()

    shift = ShiftModel(name="Night", position_id=pos.id)
    doc = DoctorModel(name="Dr. Old", email="old@hospital.org", department_id=dept.id)
    session.add(shift)
    session.add(doc)
    session.commit()

    # Existing assignment in November 2026
    old_assignment = ShiftAssignmentModel(
        doctor_id=doc.id,
        shift_id=shift.id,
        date=datetime.date(2026, 11, 15),
    )
    session.add(old_assignment)
    session.commit()

    # Draft with new assignment on Nov 1
    draft = ScheduleDraft(
        department_id=dept.id,
        target_month="2026-11",
        source_filename="new_nov.xlsx",
        status="draft",
        assignments=[
            {
                "date": "2026-11-01",
                "doctor_name": "Dr. New",
                "doctor_email": "new@hospital.org",
                "position": "ICU",
                "shift": "Night",
            }
        ],
    )
    session.add(draft)
    session.commit()

    publish_schedule_draft(session=session, draft_id=draft.id, department_id=dept.id)

    session.refresh(old_assignment)
    assert old_assignment.is_deleted is True

    active_assignments = session.query(ShiftAssignmentModel).filter(
        ShiftAssignmentModel.is_deleted == False
    ).all()
    assert len(active_assignments) == 1
    assert str(active_assignments[0].date) == "2026-11-01"



