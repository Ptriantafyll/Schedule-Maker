"""
Excel parser module
"""
import io
import calendar
import datetime
import openpyxl
from typing import Any
from src.models import Department, Doctor, Position, Shift, Team

DAY_MAP = {
    "mon": 0, "monday": 0,
    "tue": 1, "tues": 1, "tuesday": 1,
    "wed": 2, "wednesday": 2,
    "thu": 3, "thur": 3, "thurs": 3, "thursday": 3,
    "fri": 4, "friday": 4,
    "sat": 5, "saturday": 5,
    "sun": 6, "sunday": 6,
}


def _read_sheet_rows(
    ws,
    sheet_name: str,
    required_columns: list[str],
) -> list[dict[str, Any]]:
    """
    Reads a worksheet into a list of row dictionaries, validating required columns.
    """
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        raise ValueError(f"The '{sheet_name}' sheet is empty.")

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    for col in required_columns:
        if col not in headers:
            raise ValueError(
                f"Missing required column in {sheet_name} sheet: '{col}'"
            )

    rows = []
    for row in rows_iter:
        if not any(row):
            continue
        row_dict = {
            headers[i]: row[i]
            for i in range(len(headers)) if i < len(row)
        }
        rows.append(row_dict)

    return rows


def _create_default_positions(positions_in_doctors: set[str]) -> dict[str, Position]:
    """
    Creates default positions when the Shifts sheet is omitted.

    Each position receives a standard 24/7 configuration:
    a single shift named "Duty" requiring 1 doctor, with day-off
    grant, covering all 7 days of the week (Monday through Sunday).

    Args:
        positions_in_doctors: Set of unique position names found in the Doctors sheet.

    Returns:
        A dictionary mapping position names to their default Position domain objects.
    """
    pos_dict = {}
    for position_name in positions_in_doctors:
        default_shift = Shift(
            name="Duty",
            doctors_per_shift=1,
            grants_day_off=True
        )

        default_position = Position(
            name=position_name,
            shifts=[default_shift],
            duty_days={0, 1, 2, 3, 4, 5, 6}
        )
        pos_dict[position_name] = default_position

    return pos_dict


def _parse_shifts_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    positions_in_doctors: set[str],
) -> dict[str, Position]:
    """
    Parses the 'Shifts' worksheet and builds Position models with their configured shifts.

    Args:
        ws: The openpyxl worksheet for the 'Shifts' sheet.
        positions_in_doctors: Set of unique position names found in the Doctors sheet.

    Returns:
        A dictionary mapping position names to their configured Position domain objects.

    Raises:
        ValueError: If a required column is missing, a shift references an unknown position,
                    or a shift has an invalid doctors_per_shift count.
    """
    shift_rows: list[dict[str, Any]]
    shift_rows = _read_sheet_rows(
        ws,
        sheet_name="Shifts",
        required_columns=["Position", "Shift Name"]
    )

    positions: dict[str, Position] = {}
    for row in shift_rows:
        pos_name = str(row.get("Position", "")).strip()
        if pos_name not in positions_in_doctors:
            raise ValueError(
                f"Position '{pos_name}' in Shifts sheet not found in Doctors sheet."
            )

        raw_count = row.get("Doctors Per Shift")
        if raw_count is None or str(raw_count).strip() == "":
            doc_count = 1
        else:
            try:
                doc_count = int(raw_count)
            except (ValueError, TypeError) as exc:
                raise ValueError(
                    f"Invalid Doctors Per Shift '{raw_count}': must be an integer"
                ) from exc

            if doc_count <= 0:
                raise ValueError(
                    f"Doctors Per Shift must be at least 1, got '{raw_count}'"
                )

        if str(row.get("Grants Day Off", "")).strip().lower() in ("yes", "true", "1"):
            grants_day_off = True
        else:
            grants_day_off = False

        raw_shift_name = row.get("Shift Name")
        shift_name = str(raw_shift_name).strip(
        ) if raw_shift_name is not None else ""

        if not shift_name:
            raise ValueError(
                f"Shift Name cannot be empty in position '{pos_name}'")

        duty_days = _parse_duty_days(row.get("Duty Days"))
        shift = Shift(
            name=shift_name,
            doctors_per_shift=doc_count,
            grants_day_off=grants_day_off,
        )

        if pos_name not in positions:
            positions[pos_name] = Position(
                name=pos_name,
                shifts=[shift],
                duty_days=duty_days,
            )
        else:
            positions[pos_name].shifts.append(shift)
            positions[pos_name].duty_days.update(duty_days)

    for p_name in positions_in_doctors:
        if p_name not in positions:
            positions[p_name] = _create_default_positions(
                {p_name}
            )[p_name]

    return positions


def _parse_duty_days(raw_val: str | None) -> set[int]:
    """
    Parses a duty days string into a set of weekday integers (0=Monday, 6=Sunday).

    Supports:
      - None, empty, or whitespace -> all 7 days: {0, 1, 2, 3, 4, 5, 6}
      - Hyphenated ranges: "Mon-Fri", "Mon-Sun"
      - Comma-separated lists: "Mon, Wed, Fri" or "Sat, Sun"
    """
    if raw_val is None:
        return {0, 1, 2, 3, 4, 5, 6}

    cleaned = str(raw_val).strip()
    if not cleaned:
        return {0, 1, 2, 3, 4, 5, 6}

    # Handle ranges like "Mon-Fri" or "Mon-Sun"
    if "-" in cleaned:
        parts = [p.strip().lower() for p in cleaned.split("-")]
        if len(parts) != 2 or parts[0] not in DAY_MAP or parts[1] not in DAY_MAP:
            raise ValueError(f"Invalid duty days range '{raw_val}'")

        start = DAY_MAP[parts[0]]
        end = DAY_MAP[parts[1]]

        if start <= end:
            return set(range(start, end+1))

        # Handle wrap-around if someone specifies "Fri-Mon" (Fri, Sat, Sun, Mon)
        return set(range(start, 7)) | set(range(0, end+1))

    # Handle comma-separated list or single day: "Mon, Wed, Fri"
    tokens = [t.strip().lower() for t in cleaned.split(",") if t.strip()]
    if not tokens:
        return {0, 1, 2, 3, 4, 5, 6}

    result = set()
    for token in tokens:
        if token not in DAY_MAP:
            raise ValueError(f"Invalid duty day '{token}' in '{raw_val}'")
        result.add(DAY_MAP[token])

    return result


def _parse_unavailability(
    raw_val: str | None,
    year: int,
    month: int,
    days_in_month: int,
) -> set[datetime.date]:
    """
    Parses a comma-separated string of day numbers into a set of date objects.

    Args:
        raw_val: The raw cell string from Excel (e.g., "1, 2, 15", or None/empty).
        year: The target schedule year (e.g., 2026).
        month: The target schedule month (1-12).
        days_in_month: Total number of days in the month (e.g., 30 for November).

    Returns:
        A set of datetime.date objects representing the doctor's unavailable dates.

    Raises:
        ValueError: If any day token is non-numeric or falls outside 1..days_in_month.
    """
    if raw_val is None:
        return set()

    cleaned = str(raw_val).strip()
    if not cleaned:
        return set()

    tokens = [t.strip() for t in cleaned.split(",") if t.strip()]

    result: set[datetime.date] = set()
    for token in tokens:
        try:
            day = int(token)
        except ValueError as exc:
            raise ValueError(f"Invalid day: '{token}'") from exc

        if not 1 <= day <= days_in_month:
            raise ValueError(f"Day {day} out of range for {year}-{month:02d}")

        result.add(datetime.date(year, month, day))

    return result


def _parse_pre_assignments(
    raw_val: str | None,
    year: int,
    month: int,
    days_in_month: int,
    position: Position,
    has_explicit_shifts: bool,
) -> list[tuple[datetime.date, Shift]]:
    """
    Parses a doctor's pre-assigned duties from an Excel cell string.

    Accepts comma-separated "Day:ShiftName" pairs (e.g., "5:Night, 12:Morning").
    Resolves the day number to a concrete date in the target month and matches
    the shift against the doctor's position.

    Args:
        raw_val: The raw cell string from Excel (e.g., "5:Night", or None/empty).
        year: The target schedule year (e.g., 2026).
        month: The target schedule month (1-12).
        days_in_month: Total number of days in the target month.
        position: The doctor's Position domain model containing available shifts.
        has_explicit_shifts: True if the workbook included an explicit 'Shifts' sheet.

    Returns:
        A list of (datetime.date, Shift) tuples for the doctor's pre-assignments.

    Raises:
        ValueError: If a day number is invalid or out of range, if the token
                    is malformed (missing ':'), or if a shift name does not exist
                    when explicit shifts were defined.
    """
    if raw_val is None:
        return []

    cleaned = str(raw_val).strip()
    if not cleaned:
        return []

    tokens = [t.strip() for t in cleaned.split(",") if t.strip()]
    pre_assignments: list[tuple[datetime.date, Shift]] = []
    for token in tokens:
        if ":" not in token:
            raise ValueError(
                f"Invalid pre-assignment format: '{token}', expected 'Day:Shift'"
            )
        day_and_shift = [t.strip() for t in token.split(":")]
        day = day_and_shift[0]
        shift_name = str(day_and_shift[1])

        try:
            day = int(day)
        except ValueError as exc:
            raise ValueError(f"Invalid day: '{token}'") from exc

        if not 1 <= day <= days_in_month:
            raise ValueError(f"Day {day} out of range for {year}-{month:02d}")

        date = datetime.date(year, month, day)

        matching_shift = next(
            (s for s in position.shifts if s.name.lower() == shift_name.lower()),
            None
        )

        if matching_shift is None:
            if has_explicit_shifts:
                # Strict validation: admin gave a Shifts sheet, but this shift isn't in it!
                raise ValueError(
                    f"Shift '{shift_name}' not defined for position '{position.name}' in Shifts sheet."
                )
            else:
                # Flexible fallback mode: only Doctors sheet was uploaded
                # If the position has our default single "Duty" shift, rename it to what the user typed:
                if len(position.shifts) == 1 and position.shifts[0].name == "Duty":
                    position.shifts[0].name = shift_name
                    matching_shift = position.shifts[0]
                else:
                    # Otherwise create a new Shift and register it in position.shifts
                    matching_shift = Shift(
                        name=shift_name,
                        doctors_per_shift=1,
                        grants_day_off=True,
                    )
                    position.shifts.append(matching_shift)

        pre_assignments.append((date, matching_shift))

    return pre_assignments


def parse_excel_schedule_workbook(
    file_bytes: bytes,
    department_name: str,
    target_month: str,
) -> Department:
    """
    Parses an uploaded Excel schedule workbook into a Department domain model.
    """
    # 1. Parse target month
    parts = target_month.split("-")
    year, month = int(parts[0]), int(parts[1])
    days_in_month = calendar.monthrange(year, month)[1]

    # 2. Load Excel workbook from bytes
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        raise ValueError(f"Invalid or corrupted Excel file: {exc}") from exc

    # 3. Check for Doctors sheet
    if "Doctors" not in wb.sheetnames:
        raise ValueError("Workbook must contain a 'Doctors' sheet")

    # 4. Read Doctors rows & validate headers
    doctor_rows = _read_sheet_rows(
        wb["Doctors"],
        sheet_name="Doctors",
        required_columns=["Name", "Email", "Position"]
    )
    if not doctor_rows:
        raise ValueError("No doctor records found in 'Doctors' sheet")

    # 5. Validate Teams (all-or-nothing)
    has_teams = [
        bool(r.get("Team") and str(r.get("Team").strip()))
        for r in doctor_rows
    ]

    all_have_teams = all(has_teams)
    none_have_teams = not any(has_teams)

    if not (all_have_teams or none_have_teams):
        raise ValueError(
            "Team column must either be completely filled or completely blank."
        )

    # 6. Build Positions (Shifts sheet vs fallback)
    unique_positions = {
        str(r.get("Position")).strip()
        for r in doctor_rows
        if r.get("Position")
    }

    if "Shifts" in wb.sheetnames:
        positions_by_name = _parse_shifts_sheet(wb["Shifts"], unique_positions)
    else:
        positions_by_name = _create_default_positions(unique_positions)

    # 7. Build Doctor objects, teams, and return Department
    has_explicit_shifts = "Shifts" in wb.sheetnames
    teams_dict: dict[str, list[Doctor]] = {}
    all_doctors: list[Doctor] = []

    for r in doctor_rows:
        pos_name = str(r.get("Position")).strip()
        position = positions_by_name[pos_name]

        unavailability = _parse_unavailability(
            r.get("Unavailability"),
            year,
            month,
            days_in_month
        )

        pre_assignments = _parse_pre_assignments(
            r.get("Pre-Assignments"),
            year,
            month,
            days_in_month,
            position=position,
            has_explicit_shifts=has_explicit_shifts,
        )

        doc = Doctor(
            name=str(r.get("Name")).strip(),
            email=str(r.get("Email")).strip(),
            unavailability=unavailability,
            pre_assignments=pre_assignments,
        )

        position.eligible_doctors.append(doc)
        all_doctors.append(doc)
        # 2. Group into teams
        team_name = f"Team {department_name}" if none_have_teams else str(
            r.get("Team")).strip()
        if team_name not in teams_dict:
            teams_dict[team_name] = []
        teams_dict[team_name].append(doc)

    teams = [
        Team(name=name, doctors=docs)
        for name, docs in teams_dict.items()
    ]
    return Department(
        name=department_name,
        positions=list(positions_by_name.values()),
        teams=teams,
        doctor_order=all_doctors,
    )
